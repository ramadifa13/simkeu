"""Membuat template Excel migrasi data pinjaman lama."""
from pathlib import Path
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path(__file__).resolve().parents[1] / "Template_Migrasi_Reset_Total.xlsx"
MAX_ROWS = 1000

wb = Workbook()
petunjuk = wb.active
petunjuk.title = "Petunjuk"

navy = "17365D"
blue = "1F4E78"
light_blue = "D9EAF7"
yellow = "FFF2CC"
green = "E2F0D9"
red = "FCE4D6"
white = "FFFFFF"
thin = Side(style="thin", color="B7C9D6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def title(ws, text, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.font = Font(color=white, bold=True, size=14)
    cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 27

def table_header(ws, row, headers):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = f"A{row + 1}"
    ws.auto_filter.ref = f"A{row}:{chr(64 + len(headers))}{MAX_ROWS}"

def data_grid(ws, start_row, columns, date_columns=(), currency_columns=()):
    for row in range(start_row, MAX_ROWS + 1):
        for col in range(1, columns + 1):
            cell = ws.cell(row, col)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col in date_columns:
                cell.number_format = "dd/mm/yyyy"
            if col in currency_columns:
                cell.number_format = 'Rp #,##0;[Red]-Rp #,##0'

def dropdown(ws, col, formula):
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"{col}3:{col}{MAX_ROWS}")

def isi_contoh(ws, row, values):
    for col, value in enumerate(values, 1):
        ws.cell(row, col, value)

# Sheet Petunjuk
title(petunjuk, "TEMPLATE MIGRASI DATA PINJAMAN LAMA", 4)
instructions = [
    ("Tujuan", "Isi file ini untuk membangun ulang aplikasi setelah RESET TOTAL: data master, konfigurasi, riwayat pinjaman, kas, serta pendanaan Dapin."),
    ("Akun login", "Akun Admin yang melakukan reset tetap dipertahankan. Jangan tulis password atau data akun login pada file Excel ini."),
    ("Urutan pengisian", "1. Konfigurasi_Koperasi → Master_Produk → Master_Simpanan → Master_Denda → Pendanaan_Dapin → Nasabah → Pinjaman → Angsuran_Historis → Alokasi_Dapin → Saldo_Awal_Kas."),
    ("Tanggal", "Wajib tanggal asli dengan format dd/mm/yyyy. Jangan memakai angka serial Excel seperti 46304."),
    ("Master produk", "Setiap Kode Produk pada sheet Pinjaman harus tersedia pada Master_Produk. Produk menentukan plafon, tenor, bunga, frekuensi, dan aturan jaminan."),
    ("Pinjaman ulang", "Satu kali pencairan = satu baris Pinjaman baru, walaupun nasabahnya sama atau sebelumnya sudah lunas."),
    ("Angsuran", "Satu kali uang diterima = satu baris. Isi tanggal bayar dan nominal sebenarnya, termasuk denda bila ada."),
    ("Pelunasan", "Jika pelunasan terjadi sekaligus, buat satu baris pada Angsuran_Historis dengan Jenis Pembayaran = Pelunasan dan nominal uang yang benar-benar diterima."),
    ("Dapin", "Catat setiap sumber pendanaan Dapin dan alokasikan ke setiap pinjaman yang memakai modal tersebut. Ini diperlukan untuk menghitung laba bersih, bukan hanya bunga nasabah."),
    ("Saldo awal kas", "Pilih satu tanggal mulai aplikasi. Isi uang kas nyata pada tanggal itu dan jelaskan dasar perhitungannya. Jangan diisi bila seluruh mutasi historis sudah dicatat lengkap."),
    ("Nominal", "Isi nominal rupiah penuh. Contoh: 112500, bukan 112.5 dan bukan Rp112.500."),
    ("Data kosong", "Kolom kuning wajib diisi. Kolom hijau bersifat opsional tetapi sangat membantu validasi dan rekonsiliasi."),
]
petunjuk.column_dimensions["A"].width = 22
petunjuk.column_dimensions["B"].width = 110
for row, (label, text) in enumerate(instructions, 3):
    petunjuk.cell(row, 1, label).font = Font(bold=True)
    petunjuk.cell(row, 1).fill = PatternFill("solid", fgColor=light_blue)
    petunjuk.cell(row, 2, text).alignment = Alignment(wrap_text=True, vertical="top")
    petunjuk.cell(row, 1).border = petunjuk.cell(row, 2).border = border
    petunjuk.row_dimensions[row].height = 35

# Sheet Nasabah
ws = wb.create_sheet("Nasabah")
nasabah_headers = ["ID Nasabah Lama*", "NIK*", "Nama Lengkap*", "Alamat*", "No. HP", "Status", "Catatan"]
title(ws, "DATA NASABAH", len(nasabah_headers)); table_header(ws, 2, nasabah_headers)
data_grid(ws, 3, len(nasabah_headers))
for col, width in zip("ABCDEFG", [20, 22, 30, 45, 18, 15, 35]): ws.column_dimensions[col].width = width
for col in (1, 2, 3, 4):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "F", '"aktif,nonaktif"')
# Contoh lengkap: sengaja memakai ID CONTOH agar tidak dianggap data produksi.
isi_contoh(ws, 3, ["NSB-CONTOH-001", "3516070000000001", "CONTOH - Sobiroh", "Karangdiyeng, Kutorejo, Mojokerto", "081234567890", "aktif", "CONTOH PENGISIAN — ganti dengan data identitas asli."])

# Sheet Pinjaman
ws = wb.create_sheet("Pinjaman")
loan_headers = ["ID Pinjaman Lama*", "ID Nasabah Lama*", "Tanggal Cair Asli*", "Plafon (Rp)*", "Tenor (Bulan)*", "Frekuensi*", "Bunga (%)*", "Skema Bunga*", "Angsuran Normal (Rp)*", "Total Tagihan (Rp)*", "Status Saat Ini*", "Tanggal Lunas", "Catatan"]
loan_headers = ["ID Pinjaman Lama*", "ID Nasabah Lama*", "Kode Produk*", "Tanggal Cair Asli*", "Plafon (Rp)*", "Tenor (Bulan)*", "Frekuensi*", "Bunga (%)*", "Skema Bunga*", "Angsuran Normal (Rp)*", "Total Tagihan (Rp)*", "Status Saat Ini*", "Tanggal Lunas", "Catatan"]
title(ws, "DATA PENCAIRAN / PINJAMAN", len(loan_headers)); table_header(ws, 2, loan_headers)
data_grid(ws, 3, len(loan_headers), date_columns=(4, 13), currency_columns=(5, 10, 11))
for col, width in zip("ABCDEFGHIJKLMN", [20, 20, 18, 18, 18, 14, 16, 14, 16, 22, 20, 18, 18, 35]): ws.column_dimensions[col].width = width
for col in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "G", '"mingguan,dua_mingguan,bulanan"')
dropdown(ws, "I", '"flat,tahunan"')
dropdown(ws, "L", '"aktif,lunas,macet"')
isi_contoh(ws, 3, ["PJN-CONTOH-001", "NSB-CONTOH-001", "P1-2B", date(2026, 9, 12), 1000000, 2, "mingguan", 25, "flat", 156250, 1250000, "lunas", date(2026, 10, 17), "CONTOH: tanggal cair dan nilai pelunasan harus divalidasi dari catatan asli."])

# Sheet Angsuran historis
ws = wb.create_sheet("Angsuran_Historis")
payment_headers = ["ID Pembayaran", "ID Pinjaman Lama*", "Tanggal Bayar*", "Angsuran Ke", "Jenis Pembayaran*", "Pokok (Rp)", "Bunga (Rp)", "Denda (Rp)", "Total Diterima (Rp)*", "Status*", "Keterangan"]
title(ws, "RIWAYAT ANGSURAN DAN PELUNASAN", len(payment_headers)); table_header(ws, 2, payment_headers)
data_grid(ws, 3, len(payment_headers), date_columns=(3,), currency_columns=(6, 7, 8, 9))
for col, width in zip("ABCDEFGHIJK", [18, 20, 18, 14, 19, 18, 18, 18, 22, 16, 40]): ws.column_dimensions[col].width = width
for col in (2, 3, 5, 9, 10):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "E", '"cicilan,pelunasan"')
dropdown(ws, "J", '"dibayar,dibatalkan"')
contoh_bayar = [
    ["BYR-CONTOH-001", "PJN-CONTOH-001", date(2026, 9, 19), 1, "cicilan", 125000, 31250, 0, 156250, "dibayar", "CONTOH angsuran mingguan ke-1"],
    ["BYR-CONTOH-002", "PJN-CONTOH-001", date(2026, 9, 26), 2, "cicilan", 125000, 31250, 0, 156250, "dibayar", "CONTOH angsuran mingguan ke-2"],
    ["BYR-CONTOH-003", "PJN-CONTOH-001", date(2026, 10, 3), 3, "cicilan", 125000, 31250, 0, 156250, "dibayar", "CONTOH angsuran mingguan ke-3"],
    ["BYR-CONTOH-004", "PJN-CONTOH-001", date(2026, 10, 10), 4, "cicilan", 125000, 31250, 0, 156250, "dibayar", "CONTOH angsuran mingguan ke-4"],
    ["BYR-CONTOH-005", "PJN-CONTOH-001", date(2026, 10, 17), "", "pelunasan", 500000, 125000, 0, 625000, "dibayar", "CONTOH pelunasan sisa tagihan; wajib diganti nominal uang yang benar-benar diterima."],
]
for row_number, row in enumerate(contoh_bayar, 3):
    isi_contoh(ws, row_number, row)

# Sheet saldo awal
ws = wb.create_sheet("Saldo_Awal_Kas")
cash_headers = ["Tanggal Cut-off*", "Saldo Kas Fisik (Rp)*", "Mutasi Historis Sudah Diinput?*", "Dasar/Perhitungan Saldo*", "Disetujui Oleh", "Catatan"]
title(ws, "SALDO AWAL KAS / REKONSILIASI", len(cash_headers)); table_header(ws, 2, cash_headers)
data_grid(ws, 3, len(cash_headers), date_columns=(1,), currency_columns=(2,))
for col, width in zip("ABCDEF", [20, 25, 30, 60, 25, 40]): ws.column_dimensions[col].width = width
for col in (1, 2, 3, 4):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "C", '"Ya,Tidak"')
isi_contoh(ws, 3, [date(2026, 10, 17), 0, "Ya", "CONTOH: seluruh pencairan dan pembayaran sebelum tanggal cut-off sudah dicatat pada sheet terkait; masukkan kas fisik riil bila ada selisih.", "Pemilik usaha", "Jangan gunakan nilai contoh Rp0 sebagai saldo sebenarnya."])

# Konfigurasi dan master dibuat ulang setelah reset total.
ws = wb.create_sheet("Konfigurasi_Koperasi")
config_headers = ["Nama Koperasi*", "Alamat", "Telepon", "Email", "Rekening Bank", "Catatan"]
title(ws, "KONFIGURASI KOPERASI", len(config_headers)); table_header(ws, 2, config_headers)
data_grid(ws, 3, len(config_headers))
for col, width in zip("ABCDEF", [32, 55, 20, 30, 38, 45]): ws.column_dimensions[col].width = width
for col in (1,):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
isi_contoh(ws, 3, ["CONTOH KOPERASI", "Isi alamat lengkap", "081234567890", "admin@contoh.id", "Bank / Nomor Rekening", "Ganti seluruh contoh ini dengan pengaturan usaha sebenarnya."])

ws = wb.create_sheet("Master_Produk")
product_headers = ["Kode Produk*", "Nama Produk*", "Plafon Maks. (Rp)*", "Tenor Maks. (Bulan)*", "Bunga Nasabah (%)*", "Frekuensi*", "Skema Bunga*", "Perlu Jaminan?*", "Status*", "Deskripsi"]
title(ws, "MASTER PRODUK PINJAMAN", len(product_headers)); table_header(ws, 2, product_headers)
data_grid(ws, 3, len(product_headers), currency_columns=(3,))
for col, width in zip("ABCDEFGHIJ", [18, 30, 22, 20, 20, 18, 18, 20, 16, 45]): ws.column_dimensions[col].width = width
for col in range(1, 10):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "F", '"mingguan,dua_mingguan,bulanan"')
dropdown(ws, "G", '"flat,tahunan"')
dropdown(ws, "H", '"Ya,Tidak"')
dropdown(ws, "I", '"aktif,nonaktif"')
isi_contoh(ws, 3, ["P1-2B", "Pinjaman Rp1 Juta - 2 Bulan", 1000000, 2, 25, "mingguan", "flat", "Tidak", "aktif", "CONTOH dari pola angsuran Sobiroh; konfirmasi kembali produk asli."])
isi_contoh(ws, 4, ["P1-3B", "Pinjaman Rp1 Juta - 3 Bulan", 1000000, 3, 35, "mingguan", "flat", "Tidak", "aktif", "CONTOH: angsuran Rp112.500 per minggu."])
isi_contoh(ws, 5, ["P2-3B", "Pinjaman Rp2 Juta - 3 Bulan", 2000000, 3, 35, "mingguan", "flat", "Tidak", "aktif", "CONTOH: angsuran Rp225.000 per minggu."])

ws = wb.create_sheet("Master_Simpanan")
saving_headers = ["Kode Simpanan*", "Nama Simpanan*", "Minimum (Rp)", "Maksimum (Rp)", "Bagi Hasil (%)", "Status*", "Deskripsi"]
title(ws, "MASTER JENIS SIMPANAN", len(saving_headers)); table_header(ws, 2, saving_headers)
data_grid(ws, 3, len(saving_headers), currency_columns=(3, 4))
for col, width in zip("ABCDEFG", [20, 30, 20, 20, 18, 16, 45]): ws.column_dimensions[col].width = width
for col in (1, 2, 6):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "F", '"aktif,nonaktif"')
isi_contoh(ws, 3, ["ANG", "Pembayaran Angsuran", 0, 0, 0, "aktif", "WAJIB ada agar pembayaran angsuran dapat dicatat sistem."])

ws = wb.create_sheet("Master_Denda")
fine_headers = ["Plafon Pinjaman (Rp)*", "Denda per Hari (Rp)*", "Status*", "Catatan"]
title(ws, "MASTER TARIF DENDA", len(fine_headers)); table_header(ws, 2, fine_headers)
data_grid(ws, 3, len(fine_headers), currency_columns=(1, 2))
for col, width in zip("ABCD", [25, 25, 16, 50]): ws.column_dimensions[col].width = width
for col in (1, 2, 3):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "C", '"aktif,nonaktif"')
isi_contoh(ws, 3, [1000000, 1500, "aktif", "CONTOH — pastikan tarif denda yang benar."])

ws = wb.create_sheet("Pendanaan_Dapin")
dapin_headers = ["ID Dapin*", "Nama/Lembaga Dapin*", "Tanggal Dana Diterima*", "Modal Diterima (Rp)*", "Tenor (Bulan)*", "Total Wajib Kembali (Rp)*", "Tanggal Jatuh Tempo", "Status*", "Catatan"]
title(ws, "SUMBER PENDANAAN DAPIN", len(dapin_headers)); table_header(ws, 2, dapin_headers)
data_grid(ws, 3, len(dapin_headers), date_columns=(3, 7), currency_columns=(4, 6))
for col, width in zip("ABCDEFGHI", [18, 32, 22, 24, 18, 28, 22, 18, 50]): ws.column_dimensions[col].width = width
for col in (1, 2, 3, 4, 5, 6, 8):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
dropdown(ws, "H", '"aktif,lunas"')
isi_contoh(ws, 3, ["DAPIN-001", "Dapin", date(2026, 9, 12), 1000000, 3, 1060000, date(2026, 12, 12), "aktif", "CONTOH: modal Rp1 juta, wajib kembali Rp1,06 juta setelah 3 bulan."])

ws = wb.create_sheet("Alokasi_Dapin")
allocation_headers = ["ID Alokasi", "ID Pinjaman Lama*", "ID Dapin*", "Modal Dialokasikan (Rp)*", "Kewajiban Dapin (Rp)*", "Tanggal Alokasi*", "Catatan"]
title(ws, "ALOKASI MODAL DAPIN KE PINJAMAN", len(allocation_headers)); table_header(ws, 2, allocation_headers)
data_grid(ws, 3, len(allocation_headers), date_columns=(6,), currency_columns=(4, 5))
for col, width in zip("ABCDEFG", [18, 22, 18, 27, 25, 22, 55]): ws.column_dimensions[col].width = width
for col in (2, 3, 4, 5, 6):
    for row in range(3, MAX_ROWS + 1): ws.cell(row, col).fill = PatternFill("solid", fgColor=yellow)
isi_contoh(ws, 3, ["ALOK-001", "PJN-CONTOH-001", "DAPIN-001", 1000000, 1060000, date(2026, 9, 12), "CONTOH: laba bersih proyeksi = total tagihan nasabah dikurangi kewajiban Dapin."])

# Referensi
ws = wb.create_sheet("Referensi")
title(ws, "CONTOH PENGISIAN", 4)
examples = [
    ("Nasabah", "NSB-001", "Sobiroh", "Satu nasabah tetap memakai satu ID walau meminjam berulang."),
    ("Pinjaman", "PJN-001", "NSB-001", "Pencairan pertama Sobiroh; jangan gunakan ID ini untuk pinjaman berikutnya."),
    ("Angsuran", "BYR-001", "PJN-001", "Cicilan Sobiroh tanggal 19/09/2026; masukkan nominal penuh dalam rupiah."),
    ("Pelunasan", "BYR-005", "PJN-001", "Pelunasan Sobiroh tanggal 17/10/2026; nominal harus sesuai uang diterima."),
]
table_header(ws, 2, ["Sheet", "Contoh ID", "Keterkaitan", "Keterangan"])
for row, vals in enumerate(examples, 3):
    for col, val in enumerate(vals, 1):
        ws.cell(row, col, val).border = border
        ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
for col, width in zip("ABCD", [20, 20, 24, 85]): ws.column_dimensions[col].width = width

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUTPUT)
print(OUTPUT)
