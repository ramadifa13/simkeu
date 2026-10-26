from pathlib import Path
import openpyxl

src = Path(r'C:\Users\ramad\Downloads\9a3a63fa-58ae-417c-b372-27217199e370 (3).xlsx')
dst = src.with_name(src.stem + '_koreksi.xlsx')
wb = openpyxl.load_workbook(src)
ws = wb['Pinjaman']
for row in range(3, ws.max_row + 1):
    plafon = ws.cell(row, 5).value
    tenor = ws.cell(row, 6).value
    frek = str(ws.cell(row, 7).value or '').strip().lower()
    bunga = str(ws.cell(row, 8).value or '').strip()
    if plafon == 1000000 and tenor == 1 and frek == 'mingguan' and bunga == '15': code = 'SPHP 1'
    elif plafon == 1000000 and tenor == 2 and frek == 'mingguan' and bunga == '25': code = 'SPHP 2'
    elif plafon == 1000000 and tenor == 3 and frek == 'mingguan' and bunga == '35': code = 'SPHP 3'
    elif plafon == 2000000 and tenor == 1 and frek == 'mingguan' and bunga == '15': code = 'SPHP 11'
    elif plafon == 2000000 and tenor == 2 and frek == 'mingguan' and bunga == '25': code = 'SPHP 22'
    elif plafon == 2000000 and tenor == 3 and frek == 'mingguan' and bunga == '35': code = 'SPHP 33'
    elif plafon == 2000000 and tenor == 6 and frek == 'dua_mingguan' and bunga == '45': code = 'SPHP 66'
    elif plafon == 3000000 and tenor == 6 and frek == 'dua_mingguan' and bunga == '45': code = 'SPHP 666'
    else: continue
    ws.cell(row, 3).value = code

# A blank allocation means modal non-Dapin. Remove only the known placeholder row.
wa = wb['Alokasi_Dapin']
for row in range(wa.max_row, 2, -1):
    if str(wa.cell(row, 2).value or '').strip().upper() in ('PJN-CONTOH-001', 'PJN-CONTOH'):
        wa.delete_rows(row, 1)
        break

if 'Catatan_Pembersihan' in wb.sheetnames:
    del wb['Catatan_Pembersihan']
notes = wb.create_sheet('Catatan_Pembersihan')
notes.append(['Status', 'Keterangan'])
notes.append(['SESUAI', 'Alokasi_Dapin hanya diisi untuk pinjaman yang benar-benar memakai modal Dapin. Pinjaman lain dibiarkan tanpa alokasi.'])
notes.append(['WAJIB DIISI', 'Saldo Kas Fisik adalah uang kas nyata pada tanggal cut-off, bukan total pencairan atau saldo rekening Dapin.'])
notes.append(['PERLU KONFIRMASI', 'Dua produk khusus (3 juta dan 2 juta) memakai produk terdekat sesuai permintaan user; pastikan rumus angsuran historis tetap dicatat di Angsuran_Historis.'])
notes.column_dimensions['A'].width = 22
notes.column_dimensions['B'].width = 110
notes.freeze_panes = 'A2'
wb.save(dst)
print(dst)
